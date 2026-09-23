"""Importa el Excel de trayectos de FILPAC a la base de datos.

Cada fila del Excel es un movimiento del vehiculo. Entre una fila y la siguiente el
vehiculo esta quieto, asi que el DESTINO de un movimiento es el punto de inicio de la
fila siguiente (la columna 'Coordenadas de fin' de FILPAC repite la de inicio y no sirve).

Las filas de menos de UMBRAL_VIAJE_KM son maniobras menores (el vehiculo moviendose
unos metros dentro de un parqueadero): se guardan como RegistroGPS pero no crean Viaje.
"""
import re
from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation

import openpyxl
from django.db import transaction
from django.utils import timezone

from .models import RegistroGPS, ReporteGPSImportado, Vehiculo, Viaje

UMBRAL_VIAJE_KM = Decimal("0.15")


class ErrorImportacion(Exception):
    pass


def _texto(v):
    return "" if v is None else str(v).strip()


def _decimal(v):
    try:
        return Decimal(_texto(v).replace(",", "."))
    except (InvalidOperation, ValueError):
        return None


def _duracion(v):
    m = re.match(r"^\s*(\d+):(\d{2}):(\d{2})\s*$", _texto(v))
    if not m:
        return None
    h, mi, s = map(int, m.groups())
    return timedelta(hours=h, minutes=mi, seconds=s)


def _fecha_hora(v):
    if isinstance(v, datetime):
        naive = v
    else:
        try:
            naive = datetime.strptime(_texto(v)[:19], "%Y-%m-%d %H:%M:%S")
        except ValueError:
            return None
    return timezone.make_aware(naive)  # usa TIME_ZONE del proyecto (America/Bogota)


def _coords_validas(texto):
    return bool(re.match(r"^\s*-?\d+\.\d+\s+-?\d+\.\d+\s*$", _texto(texto)))


def leer_filas_excel(archivo):
    wb = openpyxl.load_workbook(archivo, data_only=True)
    ws = wb.active

    fila_enc = None
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), start=1):
        if row and row[0] == "(A) Id Objeto":
            fila_enc = i
            break
    if fila_enc is None:
        raise ErrorImportacion(
            "No parece un reporte de FILPAC: no encontre la fila de encabezados '(A) Id Objeto'."
        )

    filas = []
    for row in ws.iter_rows(min_row=fila_enc + 1, values_only=True):
        if not row or not row[0]:
            continue
        row = list(row) + [None] * (16 - len(row))
        inicio, fin = _fecha_hora(row[2]), _fecha_hora(row[4])
        if inicio is None or fin is None or not _coords_validas(row[6]):
            continue
        filas.append({
            "objeto": _texto(row[0]),
            "inicio": inicio,
            "fin": fin,
            "dir_inicio": _texto(row[3])[:300],
            "dir_fin": _texto(row[5])[:300],
            "coords_inicio": _texto(row[6]),
            "coords_fin": _texto(row[7]),
            "duracion": _duracion(row[8]),
            "distancia": _decimal(row[9]),
            "vel_max": _decimal(row[11]),
            "vel_prom": _decimal(row[12]),
            "t_mov": _duracion(row[13]),
            "t_ralenti": _duracion(row[14]),
            "t_actividad": _duracion(row[15]),
        })
    filas.sort(key=lambda f: (f["objeto"], f["inicio"]))
    return filas


@transaction.atomic
def importar_reporte(archivo, usuario):
    """Devuelve un dict con el resumen. Lanza ErrorImportacion si algo impide importar."""
    filas = leer_filas_excel(archivo)
    if not filas:
        raise ErrorImportacion("El archivo no tiene filas de trayectos validas.")

    # Vehiculos: se busca por id_objeto_gps o por placa
    objetos = sorted({f["objeto"] for f in filas})
    vehiculos, faltantes = {}, []
    for o in objetos:
        v = Vehiculo.objects.filter(id_objeto_gps=o).first() or Vehiculo.objects.filter(placa=o).first()
        if v:
            vehiculos[o] = v
        else:
            faltantes.append(o)
    if faltantes:
        raise ErrorImportacion(
            "Estos vehiculos del Excel no existen en el sistema: "
            + ", ".join(faltantes)
            + ". Crealos en el admin (Rutas > Vehiculos) con esa placa e intenta de nuevo."
        )

    if hasattr(archivo, "seek"):
        archivo.seek(0)
    reporte = ReporteGPSImportado.objects.create(
        archivo_original=archivo,
        importado_por=usuario,
        fecha_inicio_cubierta=min(f["inicio"] for f in filas).date(),
        fecha_fin_cubierta=max(f["inicio"] for f in filas).date(),
    )

    resumen = {"registros_nuevos": 0, "registros_repetidos": 0,
               "viajes_creados": 0, "maniobras_menores": 0, "viajes_sin_destino": 0}

    for objeto in objetos:
        propias = [f for f in filas if f["objeto"] == objeto]
        for i, f in enumerate(propias):
            ya_existe = RegistroGPS.objects.filter(
                objeto_gps=objeto, inicio=f["inicio"], fin=f["fin"]
            ).exists()
            if ya_existe:
                resumen["registros_repetidos"] += 1
                continue

            registro = RegistroGPS.objects.create(
                reporte_importado=reporte,
                objeto_gps=objeto,
                inicio=f["inicio"], fin=f["fin"],
                direccion_inicio=f["dir_inicio"], direccion_fin=f["dir_fin"],
                coordenadas_inicio=f["coords_inicio"], coordenadas_fin=f["coords_fin"],
                duracion=f["duracion"], distancia=f["distancia"],
                velocidad_maxima=f["vel_max"], velocidad_promedio=f["vel_prom"],
                tiempo_movimiento=f["t_mov"], tiempo_ralenti=f["t_ralenti"],
                tiempo_actividad=f["t_actividad"],
            )
            resumen["registros_nuevos"] += 1

            if f["distancia"] is None or f["distancia"] < UMBRAL_VIAJE_KM:
                resumen["maniobras_menores"] += 1
                continue

            siguiente = propias[i + 1] if i + 1 < len(propias) else None
            if siguiente is None:
                resumen["viajes_sin_destino"] += 1
            viaje = Viaje.objects.create(
                vehiculo=vehiculos[objeto],
                fecha_hora_inicio=f["inicio"], fecha_hora_fin=f["fin"],
                direccion_origen=f["dir_inicio"], direccion_destino=f["dir_fin"],
                coordenadas_inicio=f["coords_inicio"],
                coordenadas_fin=siguiente["coords_inicio"] if siguiente else "",
                distancia_total=f["distancia"], duracion_total=f["duracion"],
                velocidad_maxima=f["vel_max"], velocidad_promedio=f["vel_prom"],
                tiempo_movimiento=f["t_mov"], tiempo_ralenti=f["t_ralenti"],
                tiempo_actividad=f["t_actividad"],
            )
            registro.viaje = viaje
            registro.save(update_fields=["viaje"])
            resumen["viajes_creados"] += 1

    resumen["reporte"] = reporte
    return resumen